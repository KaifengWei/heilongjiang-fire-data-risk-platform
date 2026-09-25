"""User-facing Word report aligned with platform assessment outputs."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
BORDER = "D9D9D9"
SHADE = "F3F3F3"
BAR_FILL = "8FB5CC"
BAR_EMPTY = "EDF3F6"


def _set_run_font(run, size: float = 11, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:cs"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK


def _paragraph(doc: Document, pieces: list[tuple[str, bool]], *, size: float = 11, after: float = 8, align=None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    for text, bold in pieces:
        run = p.add_run(str(text))
        _set_run_font(run, size=size, bold=bold)


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(text)
    _set_run_font(run, size=13.5, bold=True)


def _fmt_int(value: Any) -> str:
    try:
        return f"{int(value):,}"
    except (TypeError, ValueError):
        return "—"


def _safe_float(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_border(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _summary_table(doc: Document, analysis_summary: dict[str, Any], historical_baseline: dict[str, Any], firms_intelligence: dict[str, Any]) -> None:
    summary = firms_intelligence.get("summary") or {}
    priority = historical_baseline.get("priority_regions", []) if historical_baseline else []
    priority_names = "、".join(str(x.get("region_name")) for x in priority[:3] if x.get("region_name")) or str(analysis_summary.get("main_region") or "—")
    history_status = str(historical_baseline.get("province_status") or "暂无历史同期判断") if historical_baseline and historical_baseline.get("available") else "暂无历史同期判断"
    rows = [
        ("历史同期判断", history_status),
        ("有效火点观测", _fmt_int(analysis_summary.get("active_fire_count"))),
        ("火点最集中区域", str(analysis_summary.get("main_region") or "—")),
        ("峰值日期", f"{summary.get('peak_date') or '—'}（{_fmt_int(summary.get('peak_count'))} 条）" if summary.get("peak_date") else "—"),
        ("活跃日期", f"{_fmt_int(summary.get('active_days'))} / {_fmt_int(summary.get('span_days'))}"),
        ("重点核查地区", priority_names),
    ]
    table = doc.add_table(rows=3, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    index = 0
    for row in table.rows:
        for col in (0, 2):
            label, value = rows[index]
            c1, c2 = row.cells[col], row.cells[col + 1]
            for c in (c1, c2):
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _set_cell_border(c)
            _shade(c1, SHADE)
            p = c1.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(label)
            _set_run_font(r, size=10, bold=True)
            p = c2.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            _set_run_font(r, size=10.5, bold=label in {"历史同期判断", "火点最集中区域", "峰值日期", "重点核查地区"})
            index += 1


def _top_regions(regions: list[dict[str, Any]], topn: int = 5) -> list[dict[str, Any]]:
    return sorted(regions, key=lambda x: int(x.get("observation_count") or 0), reverse=True)[:topn]


def _bar_row(cell, ratio: float, segments: int = 10) -> None:
    ratio = min(1.0, max(0.0, ratio))
    filled = max(1 if ratio > 0 else 0, round(ratio * segments))
    nested = cell.add_table(rows=1, cols=segments)
    nested.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, segment in enumerate(nested.rows[0].cells):
        _shade(segment, BAR_FILL if index < filled else BAR_EMPTY)
        _set_cell_border(segment, "FFFFFF")
        segment.width = Cm(0.55)
        p = segment.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(" ")
        _set_run_font(r, size=4)


def _visual_bar_table(doc: Document, rows: list[tuple[str, int]], caption: str) -> None:
    if not rows:
        return
    max_value = max(value for _, value in rows) or 1
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, ("项目", "数量", "相对强度")):
        _shade(cell, SHADE)
        _set_cell_border(cell)
        p = cell.paragraphs[0]
        r = p.add_run(header)
        _set_run_font(r, size=10, bold=True)
    for label, value in rows:
        cells = table.add_row().cells
        for c in cells:
            _set_cell_border(c)
        p = cells[0].paragraphs[0]
        r = p.add_run(label)
        _set_run_font(r, size=10, bold=True)
        p = cells[1].paragraphs[0]
        r = p.add_run(_fmt_int(value))
        _set_run_font(r, size=10)
        _bar_row(cells[2], value / max_value)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(caption)
    _set_run_font(r, size=9.5)


def _overall_paragraph(analysis_summary: dict[str, Any], firms_intelligence: dict[str, Any]) -> list[tuple[str, bool]]:
    total = int(analysis_summary.get("active_fire_count") or 0)
    top = _top_regions(firms_intelligence.get("regions") or [], 3)
    pieces: list[tuple[str, bool]] = [
        ("本次分析覆盖 ", False), (str(analysis_summary.get("period") or "—"), True),
        ("，共识别 ", False), (_fmt_int(total), True),
        (" 条有效主动火点观测，涉及 ", False), (_fmt_int(analysis_summary.get("affected_regions")), True),
        (" 个市（地）。", False),
    ]
    if top:
        top3_total = sum(int(row.get("observation_count") or 0) for row in top)
        parts = []
        for row in top:
            count = int(row.get("observation_count") or 0)
            share = count / total * 100 if total else 0.0
            parts.append((str(row.get("region_name") or "—"), count, share))
        pieces += [
            ("从空间分布看，", False), (parts[0][0], True), (" 最为集中，共 ", False), (_fmt_int(parts[0][1]), True),
            (" 条，约占全省 ", False), (f"{parts[0][2]:.1f}%", True), ("。", False),
        ]
        if len(parts) > 1:
            pieces += [("其次为 ", False), (parts[1][0], True), ("，共 ", False), (_fmt_int(parts[1][1]), True), (" 条。", False)]
        pieces += [("前三个地区合计约占本次有效观测的 ", False), (f"{top3_total / total * 100 if total else 0:.1f}%", True), ("，火点在空间上表现出较明显的集中分布。", False)]
    return pieces


def _time_paragraph(firms_intelligence: dict[str, Any]) -> list[tuple[str, bool]]:
    summary = firms_intelligence.get("summary") or {}
    anomalies = firms_intelligence.get("anomaly_days") or []
    repeated = firms_intelligence.get("repeated_cells") or []
    pieces: list[tuple[str, bool]] = [
        ("从时间变化看，本次任务共有 ", False), (_fmt_int(summary.get("active_days")), True), (" 个活跃日期。", False),
    ]
    if summary.get("peak_date"):
        pieces += [(str(summary.get("peak_date")), True), (" 是本次最突出的高值日，当天共记录 ", False), (_fmt_int(summary.get("peak_count")), True), (" 条火点观测。", False)]
    if anomalies:
        pieces += [("软件同时将 ", False), (str(anomalies[0].get("date") or "该日期"), True), (" 识别为火点明显增多日期，适合单独回到地图查看其空间分布。", False)]
    if repeated:
        row = repeated[0]
        pieces += [("持续性方面，", False), (str(row.get("region_name") or "部分区域"), True), (" 等地存在跨多个日期重复出现的活跃位置，其中最突出的约 ", False), (f"{row.get('grid_km', 5):g} km", True), (" 网格累计覆盖 ", False), (_fmt_int(row.get("active_days")), True), (" 个活跃日期。", False)]
    return pieces


def _assessment_paragraph(historical_baseline: dict[str, Any], firms_intelligence: dict[str, Any], priority_land_cover: dict[str, Any] | None) -> list[tuple[str, bool]]:
    pieces: list[tuple[str, bool]] = []
    if historical_baseline and historical_baseline.get("available"):
        status = str(historical_baseline.get("province_status") or "—")
        pieces += [("历史同期对比显示，全省本次火点活动整体处于 ", False), (status, True), (" 水平。", False)]
        priority = historical_baseline.get("priority_regions") or []
        if priority:
            names = "、".join(str(row.get("region_name") or "—") for row in priority[:3])
            pieces += [("平台列出的重点核查地区为 ", False), (names, True), ("。", False)]
            details = []
            for row in priority[:3]:
                name = str(row.get("region_name") or "—")
                attention = str(row.get("attention") or "")
                reasons = "；".join(str(x) for x in (row.get("reasons") or []) if x)
                detail = name
                if attention:
                    detail += f"（{attention}）"
                if reasons:
                    detail += f"：{reasons}"
                details.append(detail)
            if details:
                pieces += [("具体来看，" + "；".join(details) + "。", False)]
        else:
            pieces += [("当前没有单独列出的历史同期重点地区，可按火点数量和高值日期进行常规查看。", False)]
    else:
        pieces += [("当前任务没有形成可用的历史同期判断。", False)]
    if priority_land_cover and priority_land_cover.get("message"):
        pieces += [("地表背景方面，", False), (str(priority_land_cover.get("message")), True), ("。", False)]
    repeated = firms_intelligence.get("repeated_cells") or []
    if repeated:
        pieces += [("同时，", False), (str(repeated[0].get("region_name") or "部分地区"), True), (" 已出现持续活跃位置，适合与重点核查地区一起查看。", False)]
    return pieces


def _land_paragraph(land_cover_context: dict[str, Any]) -> list[tuple[str, bool]] | None:
    if not land_cover_context or not land_cover_context.get("available"):
        return None
    matched = int(land_cover_context.get("matched_observations") or 0)
    total = int(land_cover_context.get("total_observations") or 0)
    pieces: list[tuple[str, bool]] = [
        ("土地覆盖背景识别显示，", False), (str(land_cover_context.get("visual_title") or "已完成地表背景识别"), True), ("。", False),
    ]
    if total > 0 and matched > 0:
        pieces += [("本次共 ", False), (_fmt_int(matched), True), (" 条观测成功匹配地表背景，占有效观测的 ", False), (f"{matched / total * 100:.1f}%", True), ("。", False)]
    if land_cover_context.get("guidance"):
        pieces += [(str(land_cover_context.get("guidance")), False)]
    return pieces


def _forecast_action_text(tone: str) -> str:
    if tone == "alert":
        return "优先查看近期已有火点且重复活跃的县区，必要时安排重点核查。"
    if tone == "watch":
        return "保持对当前活跃县区的巡查，先看重复活跃位置和农田背景火点。"
    if tone == "easing":
        return "可维持常规查看，但已有持续活跃位置仍应保留在巡查清单中。"
    return "按当前重点区域继续查看即可。"


def _weather_paragraphs(weather_context: dict[str, Any] | None) -> list[list[tuple[str, bool]]]:
    if not weather_context or not weather_context.get("available"):
        return []
    result: list[list[tuple[str, bool]]] = []
    mode = str(weather_context.get("mode") or "")
    source = str(weather_context.get("source") or "")
    if mode == "historical":
        period = weather_context.get("period") or {}
        metrics = period.get("metrics") or {}
        temp = _safe_float(metrics.get("temperature_c"))
        humidity = _safe_float(metrics.get("humidity_percent"))
        wind = _safe_float(metrics.get("wind_mps"))
        rain = _safe_float(metrics.get("precipitation_mm"))
        result.append([
            ("本次任务的天气背景采用 ", False), (source or "NASA POWER Daily", True),
            ("。任务期平均气温约 ", False), (f"{temp:.1f}℃" if temp is not None else "—", True),
            ("，平均湿度约 ", False), (f"{humidity:.0f}%" if humidity is not None else "—", True),
            ("，平均风速约 ", False), (f"{wind:.1f} m/s" if wind is not None else "—", True),
            ("，日均降水约 ", False), (f"{rain:.1f} mm" if rain is not None else "—", True), ("。", False),
        ])
        peak = weather_context.get("peak_day") or {}
        if peak.get("date"):
            result.append([("火点增多日 ", False), (str(peak.get("date")), True), (" 的天气表现为：", False), (str(peak.get("summary") or "天气条件接近任务期平均"), True)])
        counties = weather_context.get("counties") or []
        if counties:
            pieces: list[tuple[str, bool]] = [("重点县区中，", False)]
            for i, row in enumerate(counties[:3]):
                if i:
                    pieces.append(("；", False))
                pieces += [(str(row.get("county_name") or "重点县区"), True), (" 为", False), (str(row.get("condition") or "已取得天气背景"), True)]
            pieces.append(("。这些天气信息可与火点高值日期和持续活跃位置一起查看。", False))
            result.append(pieces)
        return result
    result.append([("近期天气评估来自 ", False), (source or "NOAA GFS", True), ("，当前总体判断为 ", False), (str(weather_context.get("overall") or "未来天气条件已更新"), True), ("。", False)])
    for horizon in (weather_context.get("horizons") or [])[:3]:
        hours = horizon.get("hours")
        state = str(horizon.get("state") or "天气条件已更新")
        summary = str(horizon.get("summary") or "")
        tone = str(horizon.get("tone") or "normal")
        counties = [str(row.get("county_name")) for row in (horizon.get("counties") or []) if row.get("county_name")][:3]
        pieces = [(f"未来 {hours} 小时，", False), (state, True), ("。", False)]
        if summary:
            pieces.append((summary, False))
        pieces += [("平台建议：", False), (_forecast_action_text(tone), True)]
        if counties:
            pieces += [("重点查看 ", False), ("、".join(counties), True), ("。", False)]
        result.append(pieces)
    return result


def _final_paragraph(analysis_summary: dict[str, Any], firms_intelligence: dict[str, Any], historical_baseline: dict[str, Any]) -> list[tuple[str, bool]]:
    summary = firms_intelligence.get("summary") or {}
    priority = historical_baseline.get("priority_regions", []) if historical_baseline else []
    names = "、".join(str(row.get("region_name")) for row in priority[:3] if row.get("region_name"))
    pieces: list[tuple[str, bool]] = [
        ("综合来看，本次分析最需要抓住三项信息：", False),
        (str(analysis_summary.get("main_region") or "重点地区"), True), (" 是火点数量最集中的区域；", False),
        (str(summary.get("peak_date") or "峰值日期"), True), (" 是时间序列中最突出的高值日；历史同期判断为 ", False),
        (str(historical_baseline.get("province_status") or "暂无历史同期判断") if historical_baseline and historical_baseline.get("available") else "暂无历史同期判断", True), ("。", False),
    ]
    if names:
        pieces += [("后续核查可优先从 ", False), (names, True), (" 入手，再结合高值日期、持续活跃位置、地表背景和天气条件逐步缩小范围。", False)]
    else:
        pieces += [("后续可先从火点最集中区域入手，再结合高值日期、持续活跃位置、地表背景和天气条件逐步缩小范围。", False)]
    return pieces


def build_analysis_report_docx(*, task: dict[str, Any], analysis_summary: dict[str, Any], firms_intelligence: dict[str, Any], historical_baseline: dict[str, Any], land_cover_context: dict[str, Any], priority_land_cover: dict[str, Any] | None, weather_context: dict[str, Any] | None, daily_series: list[dict[str, Any]]) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("黑龙江省火点数据检测与风险评估平台分析报告")
    _set_run_font(r, size=16, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(f"{task.get('name') or '分析任务'}    {analysis_summary.get('period') or ''}")
    _set_run_font(r, size=10.5)

    _summary_table(doc, analysis_summary, historical_baseline, firms_intelligence)

    _heading(doc, "总体情况")
    _paragraph(doc, _overall_paragraph(analysis_summary, firms_intelligence))
    top_region_rows = [(str(row.get("region_name") or "—"), int(row.get("observation_count") or 0)) for row in _top_regions(firms_intelligence.get("regions") or [], 5)]
    _visual_bar_table(doc, top_region_rows, "图1 重点地区火点观测数量对比")

    _heading(doc, "时间变化")
    _paragraph(doc, _time_paragraph(firms_intelligence))
    top_dates = sorted(daily_series, key=lambda row: int(row.get("active_fire_observation_count") or 0), reverse=True)[:5]
    top_date_rows = [(str(row.get("date") or "—"), int(row.get("active_fire_observation_count") or 0)) for row in top_dates]
    _visual_bar_table(doc, top_date_rows, "图2 火点高值日期对比")

    _heading(doc, "当前评估与重点核查")
    _paragraph(doc, _assessment_paragraph(historical_baseline, firms_intelligence, priority_land_cover))

    land = _land_paragraph(land_cover_context)
    if land:
        _heading(doc, "地表背景")
        _paragraph(doc, land)

    weather_blocks = _weather_paragraphs(weather_context)
    if weather_blocks:
        _heading(doc, "天气情况")
        for block in weather_blocks:
            _paragraph(doc, block)

    _heading(doc, "综合分析")
    _paragraph(doc, _final_paragraph(analysis_summary, firms_intelligence, historical_baseline))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("生成时间：" + datetime.now().strftime("%Y-%m-%d %H:%M"))
    _set_run_font(r, size=10)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
